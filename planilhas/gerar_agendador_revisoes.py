#!/usr/bin/env python3
"""Gera o "Agendador de Revisões" em Excel (.xlsx).

Monta um calendário de estudos com revisões espaçadas automáticas, nas visões
Mês / Semana / Tabela + Painel, no estilo do calendário do app de mentoria.

Uso:
    python gerar_agendador_revisoes.py            # gera as duas versões
Saídas (na mesma pasta deste script):
    Agendador_de_Revisoes.xlsx            -> com os temas do print como exemplo
    Agendador_de_Revisoes_em_branco.xlsx  -> pronto para começar do zero

Depois de gerar, rode o recalc do LibreOffice para gravar os valores calculados.
"""

import datetime as dt
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink

# ----------------------------------------------------------------------------
# Estrutura (capacidades). Mudar aqui exige regerar o arquivo.
# ----------------------------------------------------------------------------
MAX_T = 300  # temas
NE = 6  # etapas por tema (aula + 5 revisões)
MAX_A = 200  # eventos avulsos
NB = MAX_T * NE + MAX_A  # linhas da aba Base

T0, T1 = 6, 6 + MAX_T - 1  # linhas de dados em Temas
A0, A1 = 6, 6 + MAX_A - 1  # linhas de dados em Avulsos
B0, B1 = 5, 5 + NB - 1  # linhas de dados em Base
TB0, TB1 = 6, 6 + NB - 1  # linhas de dados em Tabela

CAL_SLOTS = 6  # itens visíveis por dia no Calendário
SEM_SLOTS = 15  # itens visíveis por dia na Semana

# Intervalos de referência da Base (usados em quase todas as fórmulas)
BE = f"Base!$E${B0}:$E${B1}"  # data
BF = f"Base!$F${B0}:$F${B1}"  # tema
BG = f"Base!$G${B0}:$G${B1}"  # etapa
BH = f"Base!$H${B0}:$H${B1}"  # sigla
BI = f"Base!$I${B0}:$I${B1}"  # categoria
BJ = f"Base!$J${B0}:$J${B1}"  # área
BK = f"Base!$K${B0}:$K${B1}"  # feito
BL = f"Base!$L${B0}:$L${B1}"  # status
BM = f"Base!$M${B0}:$M${B1}"  # rótulo
BN = f"Base!$N${B0}:$N${B1}"  # código de cor
BD = f"Base!$D${B0}:$D${B1}"  # nº etapa
BP = f"Base!$P${B0}:$P${B1}"  # chave dia
BQ = f"Base!$Q${B0}:$Q${B1}"  # chave atraso

# ----------------------------------------------------------------------------
# Paleta (categorias) — barra / fundo / texto / texto concluído
# ----------------------------------------------------------------------------
CATS = [
    # nome,      código, sigla, barra,    fundo,    texto,    texto feito
    ("Teórica", "TEO", "T", "15803D", "DCFCE7", "14532D", "6B8F78"),
    ("Revisão", "REV", "R", "B45309", "FEF3C7", "78350F", "A3835F"),
    ("Prática", "PRA", "Q", "2563EB", "DBEAFE", "1E3A8A", "7189C4"),
    ("Simulado", "SIM", "SIM", "A21CAF", "FAE8FF", "701A75", "A77AAB"),
    ("Outro", "OUT", "•", "6B7280", "F3F4F6", "374151", "9CA3AF"),
]
CAT_DESC = {
    "Teórica": "Aula teórica (1º contato com o tema)",
    "Revisão": "Revisão de resumo / flashcards",
    "Prática": "Bloco de questões do tema",
    "Simulado": "Prova / simulado",
    "Outro": "Qualquer outro compromisso",
}

ETAPAS = [
    # nome,          sigla, categoria, intervalo (dias), ativa
    ("Aula teórica", "T", "Teórica", 0, "Sim"),
    ("1ª Revisão", "R1", "Revisão", 1, "Sim"),
    ("2ª Revisão", "R2", "Prática", 7, "Sim"),
    ("3ª Revisão", "R3", "Prática", 15, "Sim"),
    ("4ª Revisão", "R4", "Prática", 30, "Sim"),
    ("5ª Revisão", "R5", "Revisão", 60, "Sim"),
]

MESES = ["Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho", "Julho",
         "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro"]
DIAS = ["Domingo", "Segunda", "Terça", "Quarta", "Quinta", "Sexta", "Sábado"]
AREAS = ["Clínica Médica", "Cirurgia", "Pediatria", "Ginecologia e Obstetrícia",
         "Medicina Preventiva", "Outra", "", ""]

# Exemplo: aulas teóricas visíveis no print (set/out 2026).
# (tema, área, data da aula, etapas concluídas)
EXEMPLO_TEMAS = [
    ("Cardiopediatria", "Pediatria", "2026-09-27", {1, 2}),
    ("Nefrologia Pediátrica / ITU", "Pediatria", "2026-09-27", {1, 2}),
    ("Asma e DPOC", "Clínica Médica", "2026-09-27", {1, 2}),
    ("Doenças da Tireoide", "Clínica Médica", "2026-09-28", {1, 2}),
    ("Valvopatias", "Clínica Médica", "2026-09-28", set()),
    ("Demência e Parkinson", "Clínica Médica", "2026-09-28", set()),
    ("ATB", "Clínica Médica", "2026-10-02", {1}),
    ("Transtornos Psiquiátricos I", "Clínica Médica", "2026-10-05", set()),
    ("Síndromes Febris — Dengue", "Clínica Médica", "2026-10-08", {1}),
    ("Endocardite Infecciosa / Meningites", "Clínica Médica", "2026-10-11", set()),
    ("Especialidade Cirúrgica I", "Cirurgia", "2026-10-12", set()),
    ("Especialidade Cirúrgica II", "Cirurgia", "2026-10-12", set()),
    ("Distúrbios Hipertensivos na Gestação", "Ginecologia e Obstetrícia", "2026-10-14", set()),
    ("Sofrimento Fetal e Vitalidade Fetal", "Ginecologia e Obstetrícia", "2026-10-14", set()),
    ("Doenças da Suprarrenal", "Clínica Médica", "2026-10-16", set()),
    ("Infecções Sexualmente Transmissíveis", "Ginecologia e Obstetrícia", "2026-10-16", set()),
    ("Parasitoses Intestinais", "Pediatria", "2026-10-16", set()),
    ("HIV/AIDS", "Clínica Médica", "2026-10-16", set()),
    ("Dermatologia I e II", "Clínica Médica", "2026-10-19", set()),
    ("Bônus Oftalmo", "Cirurgia", "2026-10-20", set()),
    ("Insuficiência Cardíaca", "Clínica Médica", "2026-10-20", set()),
    ("Pneumonias Comunitárias", "Clínica Médica", "2026-10-23", set()),
    ("Tuberculose", "Clínica Médica", "2026-10-23", set()),
    ("Câncer de Pulmão e TEP", "Clínica Médica", "2026-10-23", set()),
    ("Ortopedia I e II", "Cirurgia", "2026-10-26", set()),
    ("Fraqueza Muscular — Polineuropatias", "Clínica Médica", "2026-10-27", set()),
]
EXEMPLO_AVULSOS = [
    ("2026-10-06", "UFCSPA 2026", "Simulado", "", "✓", "Prova na íntegra"),
]

# ----------------------------------------------------------------------------
# Estilos
# ----------------------------------------------------------------------------
FONT = "Arial"
SERIF = "Georgia"
INK = "1C1917"
INK2 = "57534E"
MUTED = "78716C"
LINE = "D6D3D1"
SOFT = "FAFAF9"
GREEN_DARK = "14532D"
INPUT_FILL = "FFF8DB"
INPUT_INK = "1D4ED8"


def f(size=9, bold=False, color=INK, italic=False, name=FONT, strike=False, underline=None):
    return Font(name=name, size=size, bold=bold, color=color, italic=italic,
                strike=strike, underline=underline)


def fill(c):
    return PatternFill(fill_type="solid", start_color=c, end_color=c)


def side(c=LINE, style="thin"):
    return Side(style=style, color=c)


AL_L = Alignment(horizontal="left", vertical="center")
AL_C = Alignment(horizontal="center", vertical="center")
AL_R = Alignment(horizontal="right", vertical="center")
AL_WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)


def put(ws, ref, value, font=None, align=None, fill_=None, fmt=None, border=None):
    c = ws[ref]
    c.value = value
    if font is not None:
        c.font = font
    if align is not None:
        c.alignment = align
    if fill_ is not None:
        c.fill = fill_
    if fmt is not None:
        c.number_format = fmt
    if border is not None:
        c.border = border
    return c


def mput(ws, rng, value, **kw):
    ws.merge_cells(rng)
    return put(ws, rng.split(":")[0], value, **kw)


def as_input(c, fmt=None, align=AL_L):
    c.fill = fill(INPUT_FILL)
    c.font = f(9, color=INPUT_INK)
    c.alignment = align
    c.border = Border(bottom=side("E7E5E4"))
    if fmt:
        c.number_format = fmt


def header_row(ws, row, cols_titles, height=22):
    for col, title in cols_titles:
        put(ws, f"{col}{row}", title, font=f(8, True, MUTED), align=AL_L,
            fill_=fill("F5F5F4"), border=Border(bottom=side("A8A29E")))
    ws.row_dimensions[row].height = height


def link(ws, ref, text, target, active=False, button=False):
    c = ws[ref]
    c.value = text
    c.hyperlink = Hyperlink(ref=ref, location=target, display=text)
    if button:
        c.font = f(9, True, "FFFFFF")
        c.fill = fill(GREEN_DARK)
    elif active:
        c.font = f(9, True, INK)
        c.fill = fill("E7F5EC")
    else:
        c.font = f(9, False, INK2, underline="single")
    c.alignment = AL_C


def nav(ws, cells):
    """Barra 'Mês | Semana | Tabela | Painel | + Novo tema' (hiperlinks internos)."""
    items = [("Mês", "'Calendário'!A1"), ("Semana", "'Semana'!A1"),
             ("Tabela", "'Tabela'!A1"), ("Painel", "'Painel'!A1"),
             ("+ Novo tema", "'Temas'!B6")]
    for (ref, merge), (text, target) in zip(cells, items):
        if merge:
            ws.merge_cells(merge)
        link(ws, ref, text, target, active=(ws.title == {
            "Mês": "Calendário", "Semana": "Semana", "Tabela": "Tabela",
            "Painel": "Painel"}.get(text)), button=text.startswith("+"))


def cf(ws, rng, formula, font=None, fill_=None, border=None):
    ws.conditional_formatting.add(
        rng, FormulaRule(formula=[formula], font=font, fill=fill_, border=border))


def category_rules(ws, bar_col, label_col, r_first, r_last):
    """Cores por categoria: barra lateral escura + rótulo em tom claro.
    Código 'TEO' = pendente, 'TEOx' = concluído (riscado)."""
    lab = f"{label_col}{r_first}:{label_col}{r_last}"
    bar = f"{bar_col}{r_first}:{bar_col}{r_last}"
    for _, code, _, c_bar, c_bg, c_ink, c_done in CATS:
        cf(ws, lab, f'${bar_col}{r_first}="{code}"', font=Font(color=c_ink), fill_=fill(c_bg))
        cf(ws, lab, f'${bar_col}{r_first}="{code}x"',
           font=Font(color=c_done, strike=True), fill_=fill(c_bg))
    for _, code, _, c_bar, *_ in CATS:
        cf(ws, bar, f'LEFT(${bar_col}{r_first},3)="{code}"',
           font=Font(color=c_bar), fill_=fill(c_bar))


def date_value(s):
    return dt.datetime.strptime(s, "%Y-%m-%d")


# ----------------------------------------------------------------------------
# Abas
# ----------------------------------------------------------------------------
def build_config(ws):
    ws.sheet_properties.tabColor = "78716C"
    ws.sheet_view.showGridLines = False
    for col, w in {"A": 2, "B": 6, "C": 32, "D": 12, "E": 14, "F": 18, "G": 10,
                   "H": 3, "I": 14, "J": 14, "K": 28}.items():
        ws.column_dimensions[col].width = w

    put(ws, "B1", "Configurações", font=f(18, True, INK, name=SERIF))
    put(ws, "B2", "Células amarelas podem ser alteradas. Todo o calendário se reorganiza sozinho.",
        font=f(9, color=MUTED, italic=True))

    # Etapas
    put(ws, "B4", "ETAPAS DE REVISÃO  ·  intervalos contados a partir da data da aula",
        font=f(9, True, GREEN_DARK))
    header_row(ws, 5, [("B", "Nº"), ("C", "Etapa"), ("D", "Sigla"), ("E", "Categoria (cor)"),
                       ("F", "Intervalo (dias)"), ("G", "Ativa?")])
    dv_cat = DataValidation(type="list", formula1="$C$15:$C$19", allow_blank=False)
    dv_sn = DataValidation(type="list", formula1='"Sim,Não"', allow_blank=False)
    dv_int = DataValidation(type="whole", operator="between", formula1="0", formula2="3650",
                            error="Use um número inteiro de dias (0 a 3650).", showErrorMessage=True)
    for dv in (dv_cat, dv_sn, dv_int):
        ws.add_data_validation(dv)
    for i, (nome, sigla, cat, dias, ativa) in enumerate(ETAPAS):
        r = 6 + i
        put(ws, f"B{r}", i + 1, font=f(9, color=MUTED), align=AL_C)
        for col, val in (("C", nome), ("D", sigla), ("E", cat), ("F", dias), ("G", ativa)):
            put(ws, f"{col}{r}", val)
            as_input(ws[f"{col}{r}"], align=AL_C if col in "DFG" else AL_L)
        dv_cat.add(f"E{r}")
        dv_sn.add(f"G{r}")
        dv_int.add(f"F{r}")
    ws["F6"].comment = Comment("A etapa 1 é a própria aula (0 dias). Padrão das revisões: "
                               "1, 7, 15, 30 e 60 dias depois da aula (revisão espaçada).", "Agendador")

    # Categorias
    put(ws, "B13", "CATEGORIAS E CORES  ·  a cor é fixa pela posição da linha",
        font=f(9, True, GREEN_DARK))
    header_row(ws, 14, [("B", "Nº"), ("C", "Categoria"), ("D", "Código"), ("E", "Sigla (avulsos)"),
                        ("F", "Cor"), ("G", "")])
    for i, (nome, code, sigla, c_bar, c_bg, c_ink, _) in enumerate(CATS):
        r = 15 + i
        put(ws, f"B{r}", i + 1, font=f(9, color=MUTED), align=AL_C)
        put(ws, f"C{r}", nome)
        as_input(ws[f"C{r}"])
        put(ws, f"D{r}", code, font=f(9, color=MUTED), align=AL_C)
        put(ws, f"E{r}", sigla)
        as_input(ws[f"E{r}"], align=AL_C)
        put(ws, f"F{r}", CAT_DESC[nome], font=f(8, color=c_ink), fill_=fill(c_bg),
            border=Border(left=Side(style="thick", color=c_bar)))
    ws["D14"].comment = Comment("Código interno usado pelas cores do calendário. Não altere.",
                                "Agendador")

    # Preferências
    put(ws, "B21", "PREFERÊNCIAS", font=f(9, True, GREEN_DARK))
    mput(ws, "B22:D22", "Mostrar sigla da etapa no calendário (T, R1…R5)?", font=f(9))
    put(ws, "E22", "Sim")
    as_input(ws["E22"], align=AL_C)
    dv_sn.add("E22")
    mput(ws, "B23:D23", "Dia de descanso (revisão que cair nele vai p/ o dia seguinte)", font=f(9))
    put(ws, "E23", "Nenhum")
    as_input(ws["E23"], align=AL_C)
    dv_rest = DataValidation(type="list", formula1="$J$6:$J$13", allow_blank=False)
    ws.add_data_validation(dv_rest)
    dv_rest.add("E23")
    mput(ws, "B24:D24", "(auxiliar) nº do dia de descanso — 0 = nenhum", font=f(8, color=MUTED, italic=True))
    put(ws, "E24", '=IFERROR(MATCH($E$23,$J$6:$J$13,0)-1,0)', font=f(8, color=MUTED), align=AL_C)

    # Estrutura
    put(ws, "B26", "ESTRUTURA  ·  não alterar (definida pelo gerador da planilha)",
        font=f(9, True, MUTED))
    for r, (lab, val) in enumerate((("Máx. de temas (linhas da aba Temas)", MAX_T),
                                    ("Nº de etapas por tema", NE),
                                    ("Máx. de eventos avulsos", MAX_A)), start=27):
        mput(ws, f"B{r}:D{r}", lab, font=f(9, color=MUTED))
        put(ws, f"E{r}", val, font=f(9, color=MUTED), align=AL_C)

    # Listas
    header_row(ws, 5, [("I", "Meses"), ("J", "Dia de descanso"), ("K", "Grandes áreas (editável)")])
    for i, m in enumerate(MESES):
        put(ws, f"I{6 + i}", m, font=f(9))
    for i, d in enumerate(["Nenhum"] + DIAS):
        put(ws, f"J{6 + i}", d, font=f(9))
    for i, a in enumerate(AREAS):
        put(ws, f"K{6 + i}", a if a else None)
        as_input(ws[f"K{6 + i}"])
    ws["K5"].comment = Comment("Até 8 áreas. As linhas vazias podem receber áreas novas; elas "
                               "aparecem na lista suspensa e no Painel.", "Agendador")


def build_temas(ws, exemplo):
    ws.sheet_properties.tabColor = "15803D"
    widths = {"A": 5, "B": 38, "C": 24, "D": 13, "Q": 11, "R": 13, "S": 12, "T": 8}
    for j in range(NE):
        widths[L(5 + 2 * j)] = 10
        widths[L(6 + 2 * j)] = 6
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.column_dimensions["T"].hidden = True

    put(ws, "A1", "Temas — cronograma de aulas", font=f(18, True, INK, name=SERIF))
    put(ws, "A2", "Preencha as colunas amarelas: Tema, Grande Área e Data da aula. As datas das "
                  "revisões são calculadas pela aba Config. Ao concluir, escolha ✓ (ou digite x) "
                  "na coluna Feito da etapa.", font=f(9, color=MUTED, italic=True))
    put(ws, "A3", "Para remarcar uma revisão, digite a nova data por cima da data calculada.",
        font=f(9, color=MUTED, italic=True))

    # Cabeçalho em dois níveis
    for j in range(NE):
        dc, fc = L(5 + 2 * j), L(6 + 2 * j)
        cr = 6 + j
        ws.merge_cells(f"{dc}4:{fc}4")
        put(ws, f"{dc}4", f'=Config!$C${cr}&" · +"&Config!$F${cr}&"d"', font=f(8, True, GREEN_DARK),
            align=AL_C, fill_=fill("ECFDF3"), border=Border(left=side(), right=side(), top=side()))
    heads = [("A", "Nº"), ("B", "Tema"), ("C", "Grande Área"), ("D", "Data da aula")]
    for j in range(NE):
        heads += [(L(5 + 2 * j), "Data"), (L(6 + 2 * j), "Feito")]
    heads += [("Q", "Progresso"), ("R", "Próxima etapa"), ("S", "Situação"), ("T", "aux")]
    header_row(ws, 5, heads)
    for col in [L(5 + 2 * j) for j in range(NE)] + [L(6 + 2 * j) for j in range(NE)] + ["Q", "R", "S"]:
        ws[f"{col}5"].alignment = AL_C
    ws["B5"].comment = Comment("Nome do tema/aula. Ex.: Cardiopediatria", "Agendador")
    ws["D5"].comment = Comment("Data da aula teórica, ex.: 27/09/2026. "
                               "As revisões contam a partir dela.", "Agendador")
    ws["F5"].comment = Comment("Escolha ✓ (ou digite x) quando concluir a etapa. "
                               "No calendário o item fica riscado.", "Agendador")

    dv_area = DataValidation(type="list", formula1="Config!$K$6:$K$13", allow_blank=True,
                             errorStyle="warning", showErrorMessage=True,
                             error="Área fora da lista (aba Config). Deseja manter mesmo assim?")
    dv_date = DataValidation(type="date", operator="between", formula1="36526", formula2="73051",
                             allow_blank=True, showErrorMessage=True,
                             error="Digite uma data válida (ex.: 27/09/2026).")
    dv_done = DataValidation(type="list", formula1='"✓,x"', allow_blank=True,
                             errorStyle="warning", showErrorMessage=True,
                             error="Use ✓ ou x para marcar como feito.")
    for dv in (dv_area, dv_date, dv_done):
        ws.add_data_validation(dv)
    dv_area.add(f"C{T0}:C{T1}")
    dv_date.add(f"D{T0}:D{T1}")
    for j in range(NE):
        dv_done.add(f"{L(6 + 2 * j)}{T0}:{L(6 + 2 * j)}{T1}")

    dates = [L(5 + 2 * j) for j in range(NE)]
    dones = [L(6 + 2 * j) for j in range(NE)]
    for r in range(T0, T1 + 1):
        put(ws, f"A{r}", r - T0 + 1, font=f(8, color="A8A29E"), align=AL_C)
        for col in "BC":
            as_input(ws[f"{col}{r}"])
        as_input(ws[f"D{r}"], fmt="dd/mm/yyyy", align=AL_C)
        for j in range(NE):
            cr = 6 + j
            dc, fc = dates[j], dones[j]
            put(ws, f"{dc}{r}",
                f'=IF(OR(NOT(ISNUMBER($D{r})),Config!$G${cr}<>"Sim"),"",'
                f'$D{r}+Config!$F${cr}+IF(AND(Config!$F${cr}>0,'
                f'WEEKDAY($D{r}+Config!$F${cr})=Config!$E$24),1,0))',
                font=f(9, color=INK2), align=AL_C, fmt="dd/mm/yy",
                border=Border(left=side("E7E5E4")))
            as_input(ws[f"{fc}{r}"], align=AL_C)
            ws[f"{fc}{r}"].font = f(10, True, "15803D")
        cnt = "COUNT(" + ",".join(f"{c}{r}" for c in dates) + ")"
        done_sum = "+".join(f'({d}{r}<>"")*({k}{r}<>"")' for d, k in zip(dates, dones))
        put(ws, f"Q{r}", f'=IF({cnt}=0,"",({done_sum})/{cnt})', font=f(8, color=INK2),
            align=AL_C, fmt="0%")
        mins = ",".join(f'IF(AND({d}{r}<>"",{k}{r}=""),{d}{r},99999)' for d, k in zip(dates, dones))
        put(ws, f"T{r}", f'=IF({cnt}=0,"",MIN({mins}))', font=f(8, color=MUTED))
        put(ws, f"R{r}", f'=IF(T{r}="","",IF(T{r}=99999,"",T{r}))', font=f(9, color=INK2),
            align=AL_C, fmt="ddd dd/mm")
        put(ws, f"S{r}",
            f'=IF(T{r}="","",IF(T{r}=99999,"Concluído",IF(T{r}<TODAY(),"Atrasado",'
            f'IF(T{r}=TODAY(),"Hoje","Em dia"))))', font=f(8, True, INK2), align=AL_C)

    # Formatação condicional: datas atrasadas / hoje / feitas
    for dc, fc in zip(dates, dones):
        rng = f"{dc}{T0}:{dc}{T1}"
        cf(ws, rng, f'AND(ISNUMBER({dc}{T0}),{fc}{T0}<>"")', font=Font(color="15803D"),
           fill_=fill("DCFCE7"))
        cf(ws, rng, f'AND(ISNUMBER({dc}{T0}),{fc}{T0}="",{dc}{T0}<TODAY())',
           font=Font(color="B91C1C", bold=True), fill_=fill("FEE2E2"))
        cf(ws, rng, f'AND(ISNUMBER({dc}{T0}),{fc}{T0}="",{dc}{T0}=TODAY())',
           font=Font(color="92400E", bold=True), fill_=fill("FEF3C7"))
    srng = f"S{T0}:S{T1}"
    cf(ws, srng, f'S{T0}="Atrasado"', font=Font(color="B91C1C"), fill_=fill("FEE2E2"))
    cf(ws, srng, f'S{T0}="Hoje"', font=Font(color="92400E"), fill_=fill("FEF3C7"))
    cf(ws, srng, f'S{T0}="Concluído"', font=Font(color="15803D"), fill_=fill("DCFCE7"))
    ws.conditional_formatting.add(
        f"Q{T0}:Q{T1}", DataBarRule(start_type="num", start_value=0, end_type="num",
                                    end_value=1, color="86EFAC"))

    if exemplo:
        for i, (tema, area, data, feitos) in enumerate(EXEMPLO_TEMAS):
            r = T0 + i
            ws[f"B{r}"].value = tema
            ws[f"C{r}"].value = area
            ws[f"D{r}"].value = date_value(data)
            for j in feitos:
                ws[f"{dones[j - 1]}{r}"].value = "✓"

    ws.freeze_panes = "E6"
    ws.auto_filter.ref = f"A5:S{T1}"


def build_avulsos(ws, exemplo):
    ws.sheet_properties.tabColor = "A21CAF"
    for col, w in {"A": 5, "B": 13, "C": 44, "D": 16, "E": 26, "F": 8, "G": 40}.items():
        ws.column_dimensions[col].width = w
    put(ws, "A1", "Avulsos — simulados, questões e compromissos", font=f(18, True, INK, name=SERIF))
    put(ws, "A2", "Eventos com data fixa que não seguem as revisões automáticas "
                  "(ex.: simulado UFCSPA, bloco de 100 questões, plantão).",
        font=f(9, color=MUTED, italic=True))
    put(ws, "A3", "A categoria define a cor no calendário. Marque ✓ (ou x) em Feito quando concluir.",
        font=f(9, color=MUTED, italic=True))
    header_row(ws, 5, [("A", "Nº"), ("B", "Data"), ("C", "Título / descrição"),
                       ("D", "Categoria"), ("E", "Grande Área (opcional)"), ("F", "Feito"),
                       ("G", "Observações")])
    ws["B5"].comment = Comment("Data do evento, ex.: 06/10/2026", "Agendador")
    ws["D5"].comment = Comment("Teórica, Revisão, Prática, Simulado ou Outro. "
                               "Em branco = Outro.", "Agendador")
    dv_date = DataValidation(type="date", operator="between", formula1="36526", formula2="73051",
                             allow_blank=True, showErrorMessage=True,
                             error="Digite uma data válida (ex.: 06/10/2026).")
    dv_cat = DataValidation(type="list", formula1="Config!$C$15:$C$19", allow_blank=True)
    dv_area = DataValidation(type="list", formula1="Config!$K$6:$K$13", allow_blank=True,
                             errorStyle="warning")
    dv_done = DataValidation(type="list", formula1='"✓,x"', allow_blank=True, errorStyle="warning")
    for dv in (dv_date, dv_cat, dv_area, dv_done):
        ws.add_data_validation(dv)
    dv_date.add(f"B{A0}:B{A1}")
    dv_cat.add(f"D{A0}:D{A1}")
    dv_area.add(f"E{A0}:E{A1}")
    dv_done.add(f"F{A0}:F{A1}")
    for r in range(A0, A1 + 1):
        put(ws, f"A{r}", r - A0 + 1, font=f(8, color="A8A29E"), align=AL_C)
        as_input(ws[f"B{r}"], fmt="dd/mm/yyyy", align=AL_C)
        for col in "CDEG":
            as_input(ws[f"{col}{r}"])
        as_input(ws[f"F{r}"], align=AL_C)
        ws[f"F{r}"].font = f(10, True, "15803D")
    if exemplo:
        for i, (data, titulo, cat, area, feito, obs) in enumerate(EXEMPLO_AVULSOS):
            r = A0 + i
            ws[f"B{r}"].value = date_value(data)
            ws[f"C{r}"].value = titulo
            ws[f"D{r}"].value = cat
            ws[f"E{r}"].value = area or None
            ws[f"F{r}"].value = feito or None
            ws[f"G{r}"].value = obs
    ws.freeze_panes = "A6"


def build_base(ws):
    ws.sheet_properties.tabColor = "A8A29E"
    widths = {"A": 6, "B": 8, "C": 6, "D": 6, "E": 11, "F": 34, "G": 14, "H": 6, "I": 10,
              "J": 18, "K": 6, "L": 10, "M": 40, "N": 7, "O": 5, "P": 11, "Q": 13}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    put(ws, "A1", "Base — motor de cálculo (não editar)", font=f(14, True, INK, name=SERIF))
    put(ws, "A2", f"Transforma cada tema em {NE} eventos (aula + revisões) e junta os avulsos. "
                  "Calendário, Semana, Tabela e Painel leem daqui. Não ordene nem edite.",
        font=f(9, color=MUTED, italic=True))
    header_row(ws, 4, [("A", "Nº"), ("B", "Origem"), ("C", "Linha"), ("D", "Etapa"),
                       ("E", "Data"), ("F", "Tema / título"), ("G", "Etapa (nome)"),
                       ("H", "Sigla"), ("I", "Categoria"), ("J", "Grande Área"), ("K", "Feito"),
                       ("L", "Status"), ("M", "Rótulo no calendário"), ("N", "Cor"),
                       ("O", "Seq."), ("P", "Chave dia"), ("Q", "Chave atraso")])
    TD = f"Temas!$E${T0}:$P${T1}"
    small = f(8, color=INK2)
    for r in range(B0, B1 + 1):
        row = {
            "A": "=ROW()-4",
            "B": f'=IF(A{r}<=Config!$E$27*Config!$E$28,"Tema","Avulso")',
            "C": f'=IF(B{r}="Tema",INT((A{r}-1)/Config!$E$28)+1,A{r}-Config!$E$27*Config!$E$28)',
            "D": f'=IF(B{r}="Tema",MOD(A{r}-1,Config!$E$28)+1,"")',
            "E": (f'=IF($B{r}="Tema",IF(ISNUMBER(INDEX({TD},$C{r},2*$D{r}-1)),'
                  f'INDEX({TD},$C{r},2*$D{r}-1),""),'
                  f'IF(ISNUMBER(INDEX(Avulsos!$B${A0}:$B${A1},$C{r})),'
                  f'INDEX(Avulsos!$B${A0}:$B${A1},$C{r}),""))'),
            "F": (f'=IF($E{r}="","",IF($B{r}="Tema",INDEX(Temas!$B${T0}:$B${T1},$C{r})&"",'
                  f'INDEX(Avulsos!$C${A0}:$C${A1},$C{r})&""))'),
            "G": f'=IF($E{r}="","",IF($B{r}="Tema",INDEX(Config!$C$6:$C$11,$D{r}),"Evento avulso"))',
            "H": (f'=IF($E{r}="","",IF($B{r}="Tema",INDEX(Config!$D$6:$D$11,$D{r}),'
                  f'IFERROR(INDEX(Config!$E$15:$E$19,MATCH($I{r},Config!$C$15:$C$19,0)),"•")))'),
            "I": (f'=IF($E{r}="","",IF($B{r}="Tema",INDEX(Config!$E$6:$E$11,$D{r}),'
                  f'IF(INDEX(Avulsos!$D${A0}:$D${A1},$C{r})&""="","Outro",'
                  f'INDEX(Avulsos!$D${A0}:$D${A1},$C{r}))))'),
            "J": (f'=IF($E{r}="","",IF($B{r}="Tema",INDEX(Temas!$C${T0}:$C${T1},$C{r})&"",'
                  f'INDEX(Avulsos!$E${A0}:$E${A1},$C{r})&""))'),
            "K": (f'=IF($E{r}="","",IF(IF($B{r}="Tema",INDEX({TD},$C{r},2*$D{r}),'
                  f'INDEX(Avulsos!$F${A0}:$F${A1},$C{r}))&""="","","✓"))'),
            "L": (f'=IF($E{r}="","",IF($K{r}="✓","Concluído",IF($E{r}<TODAY(),"Atrasado",'
                  f'IF($E{r}=TODAY(),"Hoje","Pendente"))))'),
            "M": (f'=IF($E{r}="","",IF($K{r}="✓","✓ ","")&'
                  f'IF(Config!$E$22="Sim",$H{r}&" · ","")&$F{r})'),
            "N": (f'=IF($E{r}="","",IFERROR(INDEX(Config!$D$15:$D$19,'
                  f'MATCH($I{r},Config!$C$15:$C$19,0)),"OUT")&IF($K{r}="✓","x",""))'),
            "O": f'=IF($E{r}="","",COUNTIF($E${B0}:$E{r},$E{r}))',
            "P": f'=IF($E{r}="","",$E{r}*100+$O{r})',
            "Q": f'=IF($L{r}="Atrasado",$E{r}*10000+$A{r},"")',
        }
        for col, val in row.items():
            c = ws[f"{col}{r}"]
            c.value = val
            c.font = small
        ws[f"E{r}"].number_format = "dd/mm/yyyy"
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:Q{B1}"


def legend_row(ws, row, first_date, last_date):
    """Legenda com contagem por categoria entre duas datas (células)."""
    for i, (nome, code, sigla, c_bar, *_rest) in enumerate(CATS):
        bc, lc = L(2 + 2 * i), L(3 + 2 * i)
        put(ws, f"{bc}{row}", None, fill_=fill(c_bar))
        put(ws, f"{lc}{row}",
            f'=" "&UPPER(Config!$C${15 + i})&"  ("&COUNTIFS({BI},Config!$C${15 + i},'
            f'{BE},">="&{first_date},{BE},"<="&{last_date})&")"',
            font=f(8, True, INK2), align=AL_L)


def build_calendario(ws):
    ws.sheet_properties.tabColor = GREEN_DARK
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.column_dimensions["A"].width = 1.5
    for d in range(7):
        ws.column_dimensions[L(2 + 2 * d)].width = 1.4
        ws.column_dimensions[L(3 + 2 * d)].width = 21
    ws.column_dimensions["P"].width = 2
    ws.column_dimensions["Q"].width = 16
    ws.column_dimensions["R"].width = 12
    ws.column_dimensions["Q"].hidden = True
    ws.column_dimensions["R"].hidden = True

    # Linha 1: marca + navegação
    put(ws, "B1", "AGENDADOR DE REVISÕES", font=f(9, True, "15803D"))
    ws.merge_cells("B1:E1")
    nav(ws, [("F1", "F1:G1"), ("H1", "H1:I1"), ("J1", "J1:K1"), ("L1", "L1:M1"), ("N1", "N1:O1")])
    ws.row_dimensions[1].height = 22

    # Linha 2: seletores
    mput(ws, "B2:C2", "MÊS ▸", font=f(8, True, MUTED), align=AL_R)
    mput(ws, "D2:E2", "Outubro")
    as_input(ws["D2"], align=AL_C)
    ws["D2"].font = f(10, True, INPUT_INK)
    mput(ws, "F2:G2", "ANO ▸", font=f(8, True, MUTED), align=AL_R)
    mput(ws, "H2:I2", 2026)
    as_input(ws["H2"], align=AL_C)
    ws["H2"].font = f(10, True, INPUT_INK)
    mput(ws, "J2:K2", "HOJE ▸", font=f(8, True, MUTED), align=AL_R)
    mput(ws, "L2:M2", "=TODAY()", font=f(9, True, INK2), align=AL_L, fmt="dd/mm/yyyy")
    mput(ws, "N2:O2",
         f'=COUNTIFS({BK},"✓",{BE},">="&$R$3,{BE},"<="&$R$4)&" de "&'
         f'COUNTIFS({BE},">="&$R$3,{BE},"<="&$R$4)&" concluídos no mês"',
         font=f(8, True, "15803D"), align=AL_R)
    ws["D2"].comment = Comment("Escolha o mês na lista. O ano fica ao lado.", "Agendador")
    dv_mes = DataValidation(type="list", formula1="Config!$I$6:$I$17", allow_blank=False)
    dv_ano = DataValidation(type="whole", operator="between", formula1="2000", formula2="2100",
                            showErrorMessage=True, error="Digite um ano entre 2000 e 2100.")
    ws.add_data_validation(dv_mes)
    ws.add_data_validation(dv_ano)
    dv_mes.add("D2")
    dv_ano.add("H2")
    ws.row_dimensions[2].height = 22

    # Linha 3: título do mês
    mput(ws, "B3:O3", '=$D$2&" "&$H$2', font=f(20, True, INK, name=SERIF), align=AL_C)
    ws.row_dimensions[3].height = 36

    # Linha 4: legenda (contagens do mês exibido)
    legend_row(ws, 4, "$R$3", "$R$4")
    mput(ws, "L4:O4", "riscado com ✓ = concluído", font=f(8, color=MUTED, italic=True), align=AL_R)
    ws.row_dimensions[4].height = 18
    ws.row_dimensions[5].height = 6

    # Auxiliares (ocultas)
    aux = [("Nº do mês", "=IFERROR(MATCH($D$2,Config!$I$6:$I$17,0),MONTH(TODAY()))", "0"),
           ("1º dia do mês", "=DATE(IF(ISNUMBER($H$2),$H$2,YEAR(TODAY())),$R$2,1)", "dd/mm/yyyy"),
           ("Último dia do mês", "=DATE(YEAR($R$3),$R$2+1,0)", "dd/mm/yyyy"),
           ("Início da grade (dom)", "=$R$3-WEEKDAY($R$3)+1", "dd/mm/yyyy")]
    for i, (lab, frm, fmt) in enumerate(aux):
        put(ws, f"Q{2 + i}", lab, font=f(8, color=MUTED))
        put(ws, f"R{2 + i}", frm, font=f(8, color=MUTED), fmt=fmt)

    # Linha 6: dias da semana
    for d, nome in enumerate(["DOM", "SEG", "TER", "QUA", "QUI", "SEX", "SAB"]):
        bc, lc = L(2 + 2 * d), L(3 + 2 * d)
        ws.merge_cells(f"{bc}6:{lc}6")
        put(ws, f"{bc}6", nome, font=f(8, True, MUTED), align=AL_C, fill_=fill(SOFT),
            border=Border(top=side(), bottom=side(), left=side() if d == 0 else None))
        ws[f"{lc}6"].border = Border(top=side(), bottom=side(), right=side())
    ws.row_dimensions[6].height = 22

    # Grade 6 semanas x 7 dias (cada semana = 1 linha de data + slots + 1 linha "+N mais")
    WK = CAL_SLOTS + 2
    r_first, r_last = 7, 7 + 6 * WK - 1
    for w in range(6):
        r0 = 7 + w * WK
        ws.row_dimensions[r0].height = 19
        for k in range(1, CAL_SLOTS + 1):
            ws.row_dimensions[r0 + k].height = 16.5
        ws.row_dimensions[r0 + CAL_SLOTS + 1].height = 12
        for d in range(7):
            bc, lc = L(2 + 2 * d), L(3 + 2 * d)
            dref = f"{lc}{r0}"
            put(ws, dref, f"=$R$5+{7 * w + d}", font=f(9, True, INK2), align=AL_L, fmt="d")
            for k in range(1, CAL_SLOTS + 1):
                r = r0 + k
                key = f"{dref}*100+{k}"
                put(ws, f"{bc}{r}", f'=IFERROR(INDEX({BN},MATCH({key},{BP},0)),"")',
                    font=f(6, color="FFFFFF"), align=AL_L)
                put(ws, f"{lc}{r}", f'=IFERROR(INDEX({BM},MATCH({key},{BP},0)),"")',
                    font=f(8, color=INK2), align=AL_L)
            rm = r0 + CAL_SLOTS + 1
            put(ws, f"{lc}{rm}",
                f'=IF(COUNTIF({BE},{dref})>{CAL_SLOTS},"+"&(COUNTIF({BE},{dref})-{CAL_SLOTS})'
                f'&" mais","")', font=f(7, True, MUTED, italic=True), align=AL_C)
            # bordas do bloco do dia
            for r in range(r0, rm + 1):
                b_left = side() if d == 0 else None
                ws[f"{bc}{r}"].border = Border(left=b_left, bottom=side() if r == rm else None)
                ws[f"{lc}{r}"].border = Border(right=side(), bottom=side() if r == rm else None)

    # Formatação condicional (ordem = prioridade)
    for d in range(7):
        bc, lc = L(2 + 2 * d), L(3 + 2 * d)
        cf(ws, f"{lc}{r_first}:{lc}{r_last}", f"{lc}{r_first}=TODAY()",
           font=Font(color="FFFFFF", bold=True), fill_=fill(GREEN_DARK))
    for d in range(7):
        bc, lc = L(2 + 2 * d), L(3 + 2 * d)
        category_rules(ws, bc, lc, r_first, r_last)
    for d in range(7):
        bc, lc = L(2 + 2 * d), L(3 + 2 * d)
        block = f"{bc}{r_first}:{lc}{r_last}"
        day = f"($R$5+INT((ROW({bc}{r_first})-{r_first})/{WK})*7+{d})"
        cf(ws, block, f"{day}=TODAY()", fill_=fill("EEFBF3"))
        cf(ws, block, f"MONTH({day})<>$R$2", font=Font(color="A8A29E"), fill_=fill("F5F5F4"))

    ws.freeze_panes = "A7"
    ws.print_area = f"A1:O{r_last}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.page_margins.top = ws.page_margins.bottom = 0.4


def build_semana(ws):
    ws.sheet_properties.tabColor = "2563EB"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.column_dimensions["A"].width = 1.5
    for d in range(7):
        ws.column_dimensions[L(2 + 2 * d)].width = 1.4
        ws.column_dimensions[L(3 + 2 * d)].width = 25
    ws.column_dimensions["Q"].width = 16
    ws.column_dimensions["R"].width = 12
    ws.column_dimensions["Q"].hidden = True
    ws.column_dimensions["R"].hidden = True

    put(ws, "B1", "AGENDADOR DE REVISÕES", font=f(9, True, "15803D"))
    ws.merge_cells("B1:E1")
    nav(ws, [("F1", "F1:G1"), ("H1", "H1:I1"), ("J1", "J1:K1"), ("L1", "L1:M1"), ("N1", "N1:O1")])
    ws.row_dimensions[1].height = 22

    mput(ws, "B2:C2", "SEMANA DE ▸", font=f(8, True, MUTED), align=AL_R)
    mput(ws, "D2:E2", "=TODAY()", fmt="dd/mm/yyyy")
    as_input(ws["D2"], fmt="dd/mm/yyyy", align=AL_C)
    ws["D2"].font = f(10, True, INPUT_INK)
    mput(ws, "F2:K2", "← digite qualquer data para ver a semana dela (padrão: hoje)",
         font=f(8, color=MUTED, italic=True), align=AL_L)
    ws["D2"].comment = Comment("Padrão: =HOJE(). Digite uma data por cima para ver outra semana; "
                               "para voltar, digite =HOJE() de novo.", "Agendador")
    ws.row_dimensions[2].height = 22

    meses = "Config!$I$6:$I$17"
    mput(ws, "B3:O3",
         f'=DAY($R$2)&" de "&LOWER(INDEX({meses},MONTH($R$2)))&" a "&DAY($R$3)&" de "&'
         f'LOWER(INDEX({meses},MONTH($R$3)))&" de "&YEAR($R$3)',
         font=f(20, True, INK, name=SERIF), align=AL_C)
    ws.row_dimensions[3].height = 36
    legend_row(ws, 4, "$R$2", "$R$3")
    mput(ws, "L4:O4", "contagens desta semana", font=f(8, color=MUTED, italic=True), align=AL_R)
    ws.row_dimensions[4].height = 18
    ws.row_dimensions[5].height = 6

    put(ws, "Q2", "Início (dom)", font=f(8, color=MUTED))
    put(ws, "R2", "=IF(ISNUMBER($D$2),$D$2,TODAY())-WEEKDAY(IF(ISNUMBER($D$2),$D$2,TODAY()))+1", font=f(8, color=MUTED), fmt="dd/mm/yyyy")
    put(ws, "Q3", "Fim (sáb)", font=f(8, color=MUTED))
    put(ws, "R3", "=$R$2+6", font=f(8, color=MUTED), fmt="dd/mm/yyyy")

    r_head, r_date, r_sum = 6, 7, 8
    r_s0 = 9
    r_more = r_s0 + SEM_SLOTS
    for d, nome in enumerate(["DOMINGO", "SEGUNDA", "TERÇA", "QUARTA", "QUINTA", "SEXTA", "SÁBADO"]):
        bc, lc = L(2 + 2 * d), L(3 + 2 * d)
        ws.merge_cells(f"{bc}{r_head}:{lc}{r_head}")
        put(ws, f"{bc}{r_head}", nome, font=f(8, True, MUTED), align=AL_C, fill_=fill(SOFT),
            border=Border(top=side(), bottom=side(), left=side() if d == 0 else None))
        ws[f"{lc}{r_head}"].border = Border(top=side(), bottom=side(), right=side())
        dref = f"{lc}{r_date}"
        put(ws, dref, f"=$R$2+{d}", font=f(14, True, INK2), align=AL_L, fmt="dd/mm")
        put(ws, f"{lc}{r_sum}",
            f'=IF(COUNTIF({BE},{dref})=0,"livre",COUNTIF({BE},{dref})&'
            f'IF(COUNTIF({BE},{dref})=1," item · "," itens · ")&'
            f'COUNTIFS({BE},{dref},{BK},"✓")&" ✓")', font=f(8, color=MUTED, italic=True),
            align=AL_L)
        for k in range(1, SEM_SLOTS + 1):
            r = r_s0 + k - 1
            key = f"{dref}*100+{k}"
            put(ws, f"{bc}{r}", f'=IFERROR(INDEX({BN},MATCH({key},{BP},0)),"")',
                font=f(6, color="FFFFFF"), align=AL_L)
            put(ws, f"{lc}{r}", f'=IFERROR(INDEX({BM},MATCH({key},{BP},0)),"")',
                font=f(9, color=INK2), align=AL_L)
        put(ws, f"{lc}{r_more}",
            f'=IF(COUNTIF({BE},{dref})>{SEM_SLOTS},"+"&(COUNTIF({BE},{dref})-{SEM_SLOTS})'
            f'&" mais (veja a aba Tabela)","")', font=f(7, True, MUTED, italic=True), align=AL_C)
        for r in range(r_date, r_more + 1):
            ws[f"{bc}{r}"].border = Border(left=side() if d == 0 else None,
                                           bottom=side() if r == r_more else None)
            ws[f"{lc}{r}"].border = Border(right=side(), bottom=side() if r == r_more else None)
    ws.row_dimensions[r_head].height = 22
    ws.row_dimensions[r_date].height = 24
    ws.row_dimensions[r_sum].height = 15
    for r in range(r_s0, r_more):
        ws.row_dimensions[r].height = 18
    ws.row_dimensions[r_more].height = 13

    for d in range(7):
        bc, lc = L(2 + 2 * d), L(3 + 2 * d)
        cf(ws, f"{lc}{r_date}", f"{lc}{r_date}=TODAY()",
           font=Font(color="FFFFFF", bold=True), fill_=fill(GREEN_DARK))
    for d in range(7):
        bc, lc = L(2 + 2 * d), L(3 + 2 * d)
        category_rules(ws, bc, lc, r_s0, r_more - 1)
    for d in range(7):
        bc, lc = L(2 + 2 * d), L(3 + 2 * d)
        cf(ws, f"{bc}{r_date}:{lc}{r_more}", f"($R$2+{d})=TODAY()", fill_=fill("EEFBF3"))

    ws.freeze_panes = "A7"
    ws.print_area = f"A1:O{r_more}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def build_tabela(ws):
    ws.sheet_properties.tabColor = "57534E"
    ws.sheet_view.showGridLines = False
    for col, w in {"A": 6, "B": 12, "C": 14, "D": 40, "E": 16, "F": 7, "G": 12, "H": 24,
                   "I": 12, "J": 12, "K": 8}.items():
        ws.column_dimensions[col].width = w
    ws.column_dimensions["J"].hidden = True
    ws.column_dimensions["K"].hidden = True
    nav(ws, [("D1", None), ("E1", None), ("G1", None), ("H1", None), ("I1", None)])
    put(ws, "A2", "Tabela — todos os eventos em ordem de data", font=f(18, True, INK, name=SERIF))
    put(ws, "A3", "Use os filtros do cabeçalho (ex.: Status = Atrasado, Categoria = Prática). "
                  "Não edite aqui: altere em Temas ou Avulsos.", font=f(9, color=MUTED, italic=True))
    header_row(ws, 5, [("A", "#"), ("B", "Data"), ("C", "Dia"), ("D", "Tema / título"),
                       ("E", "Etapa"), ("F", "Sigla"), ("G", "Categoria"), ("H", "Grande Área"),
                       ("I", "Status"), ("J", "chave"), ("K", "linha")])
    for r in range(TB0, TB1 + 1):
        k = f"$K{r}"
        row = {
            "J": f'=IFERROR(SMALL({BP},ROW()-{TB0 - 1}),"")',
            "K": f'=IF(J{r}="","",MATCH(J{r},{BP},0))',
            "A": f'=IF({k}="","",ROW()-{TB0 - 1})',
            "B": f'=IF({k}="","",INDEX({BE},{k}))',
            "C": f'=IF({k}="","",B{r})',
            "D": f'=IF({k}="","",INDEX({BF},{k}))',
            "E": f'=IF({k}="","",INDEX({BG},{k}))',
            "F": f'=IF({k}="","",INDEX({BH},{k}))',
            "G": f'=IF({k}="","",INDEX({BI},{k}))',
            "H": f'=IF({k}="","",INDEX({BJ},{k}))',
            "I": f'=IF({k}="","",INDEX({BL},{k}))',
        }
        for col, val in row.items():
            c = ws[f"{col}{r}"]
            c.value = val
            c.font = f(9, color=INK2 if col not in "AJK" else "A8A29E")
            c.border = Border(bottom=side("F5F5F4"))
        ws[f"B{r}"].number_format = "dd/mm/yyyy"
        ws[f"C{r}"].number_format = "dddd"
        for col in "ABFGI":
            ws[f"{col}{r}"].alignment = AL_C
        ws[f"C{r}"].alignment = AL_L
    rng_g = f"G{TB0}:G{TB1}"
    for nome, code, sigla, c_bar, c_bg, c_ink, _ in CATS:
        cf(ws, rng_g, f'$G{TB0}="{nome}"', font=Font(color=c_ink, bold=True), fill_=fill(c_bg))
    rng_i = f"I{TB0}:I{TB1}"
    cf(ws, rng_i, f'I{TB0}="Atrasado"', font=Font(color="B91C1C", bold=True), fill_=fill("FEE2E2"))
    cf(ws, rng_i, f'I{TB0}="Hoje"', font=Font(color="92400E", bold=True), fill_=fill("FEF3C7"))
    cf(ws, rng_i, f'I{TB0}="Concluído"', font=Font(color="15803D"), fill_=fill("DCFCE7"))
    cf(ws, f"D{TB0}:D{TB1}", f'$I{TB0}="Concluído"', font=Font(color="A8A29E", strike=True))
    cf(ws, f"A{TB0}:I{TB1}", f"$B{TB0}=TODAY()", fill_=fill("EEFBF3"))
    ws.freeze_panes = "A6"
    ws.auto_filter.ref = f"A5:I{TB1}"


def build_painel(ws):
    ws.sheet_properties.tabColor = "B45309"
    ws.sheet_view.showGridLines = False
    for col, w in {"A": 2, "B": 14, "C": 34, "D": 20, "E": 13, "F": 13, "G": 13, "H": 13,
                   "I": 3, "J": 14, "K": 34, "L": 20, "M": 13, "N": 13}.items():
        ws.column_dimensions[col].width = w
    nav(ws, [("J1", None), ("K1", None), ("L1", None), ("M1", None), ("N1", None)])
    put(ws, "B2", "Painel", font=f(20, True, INK, name=SERIF))
    put(ws, "B3", "Hoje é", font=f(9, color=MUTED))
    put(ws, "C3", "=TODAY()", font=f(9, True, INK2), fmt="dddd, dd/mm/yyyy", align=AL_L)
    ws.row_dimensions[2].height = 30

    def section(ref, text):
        put(ws, ref, text, font=f(9, True, GREEN_DARK))

    # KPIs
    kpis = [
        ("B", "Temas", f'=COUNTIF(Temas!$B${T0}:$B${T1},"?*")', "0"),
        ("C", "Eventos agendados", f"=COUNT({BE})", "0"),
        ("D", "Concluídos", f'=COUNTIF({BK},"✓")', "0"),
        ("E", "% concluído", "=IF(C6=0,0,D6/C6)", "0%"),
        ("F", "Atrasados", f'=COUNTIF({BL},"Atrasado")', "0"),
        ("G", "Para hoje", f"=COUNTIF({BE},TODAY())", "0"),
        ("H", "Próx. 7 dias", f'=COUNTIFS({BE},">"&TODAY(),{BE},"<="&TODAY()+7)', "0"),
    ]
    for col, lab, frm, fmt in kpis:
        put(ws, f"{col}5", lab, font=f(8, True, MUTED), align=AL_C, fill_=fill(SOFT),
            border=Border(top=side(), left=side(), right=side()))
        put(ws, f"{col}6", frm, font=f(20, True, INK, name=SERIF), align=AL_C, fmt=fmt,
            fill_=fill(SOFT), border=Border(bottom=side(), left=side(), right=side()))
    ws["F6"].font = f(20, True, "B91C1C", name=SERIF)
    ws["E6"].font = f(20, True, "15803D", name=SERIF)
    ws.row_dimensions[6].height = 36

    # Por categoria
    section("C8", "POR CATEGORIA")
    header_row(ws, 9, [("C", "Categoria"), ("D", "Agendados"), ("E", "Concluídos"),
                       ("F", "% concluído"), ("G", "Atrasados")])
    for i, (nome, code, sigla, c_bar, c_bg, c_ink, _) in enumerate(CATS):
        r = 10 + i
        put(ws, f"C{r}", f"=Config!$C${15 + i}", font=f(9, True, c_ink), fill_=fill(c_bg),
            border=Border(left=Side(style="thick", color=c_bar)))
        put(ws, f"D{r}", f"=COUNTIF({BI},C{r})", font=f(9), align=AL_C)
        put(ws, f"E{r}", f'=COUNTIFS({BI},C{r},{BK},"✓")', font=f(9), align=AL_C)
        put(ws, f"F{r}", f'=IF(D{r}=0,"",E{r}/D{r})', font=f(9), align=AL_C, fmt="0%")
        put(ws, f"G{r}", f'=COUNTIFS({BI},C{r},{BL},"Atrasado")', font=f(9, color="B91C1C"),
            align=AL_C)
    put(ws, "C15", "Total", font=f(9, True))
    for col in "DEG":
        put(ws, f"{col}15", f"=SUM({col}10:{col}14)", font=f(9, True), align=AL_C,
            border=Border(top=side("A8A29E")))
    put(ws, "F15", '=IF(D15=0,"",E15/D15)', font=f(9, True), align=AL_C, fmt="0%",
        border=Border(top=side("A8A29E")))
    ws.conditional_formatting.add("F10:F14", DataBarRule(start_type="num", start_value=0,
                                                         end_type="num", end_value=1,
                                                         color="86EFAC"))

    # Por grande área
    section("J8", "POR GRANDE ÁREA")
    header_row(ws, 9, [("J", ""), ("K", "Grande Área"), ("L", "Temas"), ("M", "Eventos"),
                       ("N", "% concluído")])
    for i in range(8):
        r = 10 + i
        put(ws, f"K{r}", f'=IF(Config!$K${6 + i}="","",Config!$K${6 + i})', font=f(9, True))
        put(ws, f"L{r}", f'=IF(K{r}="","",COUNTIF(Temas!$C${T0}:$C${T1},K{r}))', font=f(9),
            align=AL_C)
        put(ws, f"M{r}", f'=IF(K{r}="","",COUNTIF({BJ},K{r}))', font=f(9), align=AL_C)
        put(ws, f"N{r}", f'=IF(K{r}="","",IF(M{r}=0,"",COUNTIFS({BJ},K{r},{BK},"✓")/M{r}))',
            font=f(9), align=AL_C, fmt="0%")
    ws.conditional_formatting.add("N10:N17", DataBarRule(start_type="num", start_value=0,
                                                         end_type="num", end_value=1,
                                                         color="86EFAC"))

    # Para hoje
    N_LIST = 15
    section("B20", "PARA HOJE")
    header_row(ws, 21, [("B", "Sigla"), ("C", "Tema / título"), ("D", "Etapa"),
                        ("E", "Categoria"), ("F", "Status")])
    for k in range(1, N_LIST + 1):
        r = 21 + k
        m = f"MATCH(TODAY()*100+{k},{BP},0)"
        for col, src in (("B", BH), ("C", BF), ("D", BG), ("E", BI), ("F", BL)):
            put(ws, f"{col}{r}", f'=IFERROR(INDEX({src},{m}),"")', font=f(9, color=INK2),
                align=AL_C if col in "BEF" else AL_L, border=Border(bottom=side("F5F5F4")))
    put(ws, f"B{22 + N_LIST}",
        f'=IF($G$6>{N_LIST},"+"&($G$6-{N_LIST})&" itens além desta lista (veja a aba Semana)","")',
        font=f(8, color=MUTED, italic=True))
    cf(ws, f"C22:C{21 + N_LIST}", '$F22="Concluído"', font=Font(color="A8A29E", strike=True))
    rng = f"E22:E{21 + N_LIST}"
    for nome, code, sigla, c_bar, c_bg, c_ink, _ in CATS:
        cf(ws, rng, f'$E22="{nome}"', font=Font(color=c_ink, bold=True), fill_=fill(c_bg))

    # Atrasadas
    section("J20", "ATRASADAS  ·  mais antigas primeiro")
    header_row(ws, 21, [("J", "Data"), ("K", "Tema / título"), ("L", "Etapa"),
                        ("M", "Categoria"), ("N", "Dias de atraso")])
    for k in range(1, N_LIST + 1):
        r = 21 + k
        m = f"MATCH(SMALL({BQ},{k}),{BQ},0)"
        for col, src in (("J", BE), ("K", BF), ("L", BG), ("M", BI)):
            put(ws, f"{col}{r}", f'=IFERROR(INDEX({src},{m}),"")', font=f(9, color=INK2),
                align=AL_C if col in "JM" else AL_L, border=Border(bottom=side("F5F5F4")))
        ws[f"J{r}"].number_format = "dd/mm/yyyy"
        put(ws, f"N{r}", f'=IF(J{r}="","",TODAY()-J{r})', font=f(9, True, "B91C1C"),
            align=AL_C, fmt="0", border=Border(bottom=side("F5F5F4")))
    put(ws, f"J{22 + N_LIST}",
        f'=IF($F$6>{N_LIST},"+"&($F$6-{N_LIST})&" atrasadas além desta lista '
        f'(filtre Status = Atrasado na aba Tabela)","")', font=f(8, color=MUTED, italic=True))
    rng = f"M22:M{21 + N_LIST}"
    for nome, code, sigla, c_bar, c_bg, c_ink, _ in CATS:
        cf(ws, rng, f'$M22="{nome}"', font=Font(color=c_ink, bold=True), fill_=fill(c_bg))

    # Carga dos próximos 14 dias
    r_h = 40
    section(f"B{r_h - 1}", "CARGA DOS PRÓXIMOS 14 DIAS")
    header_row(ws, r_h, [("B", "Data"), ("C", "Total de itens"), ("D", "=Config!$C$15"),
                         ("E", "=Config!$C$16"), ("F", "=Config!$C$17"), ("G", "=Config!$C$18"),
                         ("H", "=Config!$C$19")])
    for col in "DEFGH":
        ws[f"{col}{r_h}"].alignment = AL_C
    for i in range(14):
        r = r_h + 1 + i
        put(ws, f"B{r}", f"=TODAY()+{i}", font=f(9, True if i == 0 else False, INK2),
            fmt="ddd dd/mm", align=AL_L, border=Border(bottom=side("F5F5F4")))
        for col in "DEFGH":
            put(ws, f"{col}{r}", f'=COUNTIFS({BE},$B{r},{BI},{col}${r_h})', font=f(9, color=INK2),
                align=AL_C, fmt="0;-0;\"·\"", border=Border(bottom=side("F5F5F4")))
        put(ws, f"C{r}", f"=SUM(D{r}:H{r})", font=f(9, True, INK), align=AL_L,
            border=Border(bottom=side("F5F5F4")))
    ws.conditional_formatting.add(
        f"C{r_h + 1}:C{r_h + 14}",
        DataBarRule(start_type="num", start_value=0, end_type="max", color="93C5FD"))

    # Progresso por etapa
    section(f"J{r_h - 1}", "PROGRESSO POR ETAPA (temas)")
    header_row(ws, r_h, [("J", "Sigla"), ("K", "Etapa"), ("L", "Agendadas"), ("M", "Concluídas"),
                         ("N", "% concluído")])
    for i in range(NE):
        r = r_h + 1 + i
        put(ws, f"J{r}", f"=Config!$D${6 + i}", font=f(9, True, INK2), align=AL_C)
        put(ws, f"K{r}", f"=Config!$C${6 + i}", font=f(9, color=INK2))
        put(ws, f"L{r}", f'=COUNTIFS({BD},{i + 1},{BE},">0")', font=f(9), align=AL_C)
        put(ws, f"M{r}", f'=COUNTIFS({BD},{i + 1},{BK},"✓")', font=f(9), align=AL_C)
        put(ws, f"N{r}", f'=IF(L{r}=0,"",M{r}/L{r})', font=f(9), align=AL_C, fmt="0%")
    ws.conditional_formatting.add(
        f"N{r_h + 1}:N{r_h + NE}",
        DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="86EFAC"))
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def build_como_usar(ws, exemplo):
    ws.sheet_properties.tabColor = "1C1917"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 4
    ws.column_dimensions["C"].width = 26
    ws.column_dimensions["D"].width = 90

    put(ws, "B2", "Agendador de Revisões", font=f(24, True, INK, name=SERIF))
    put(ws, "B3", "Monte seu calendário de estudos com revisões espaçadas automáticas — "
                  "visões Mês, Semana e Tabela, como no calendário do app.",
        font=f(10, color=MUTED, italic=True))
    r = 5

    def h(text):
        nonlocal r
        r += 1
        put(ws, f"B{r}", text, font=f(10, True, GREEN_DARK))
        ws[f"B{r}"].border = Border(bottom=side("A8A29E"))
        for col in "CD":
            ws[f"{col}{r}"].border = Border(bottom=side("A8A29E"))
        r += 1

    def item(label, text, label_font=None):
        nonlocal r
        put(ws, f"C{r}", label, font=label_font or f(9, True, INK), align=AL_WRAP)
        put(ws, f"D{r}", text, font=f(9, color=INK2), align=AL_WRAP)
        lines = max(1, len(text) // 95 + 1)
        ws.row_dimensions[r].height = 14 * lines + 4
        r += 1

    h("COMO USAR — 4 PASSOS")
    item("1. Temas", "Para cada aula, digite o Tema, a Grande Área e a Data da aula (células "
                     "amarelas). As datas da aula e das 5 revisões aparecem sozinhas.")
    item("2. Avulsos", "Lance simulados, blocos de questões e compromissos com data fixa "
                       "(ex.: simulado UFCSPA 2026).")
    item("3. Calendário", "Escolha Mês e Ano no topo. Cada dia mostra até 6 itens com a cor da "
                          "categoria; se houver mais, aparece “+N mais” (veja tudo em Semana ou "
                          "Tabela). O dia de hoje fica destacado em verde.")
    item("4. Marcar feito", "Ao concluir uma etapa, escolha ✓ (ou digite x) na coluna Feito daquela "
                            "etapa em Temas/Avulsos. No calendário o item fica riscado com ✓.")

    h("PERSONALIZAR (aba Config)")
    item("Intervalos", "Padrão: aula (0) e revisões com 1, 7, 15, 30 e 60 dias. Mude os números e "
                       "o calendário inteiro se reorganiza.")
    item("Categoria da etapa", "Define a cor no calendário: Teórica, Revisão, Prática (questões), "
                               "Simulado ou Outro.")
    item("Ativa?", "“Não” desliga a etapa — ela some do calendário, da tabela e do painel.")
    item("Dia de descanso", "Revisões que cairiam nesse dia da semana vão para o dia seguinte.")
    item("Mostrar sigla", "Exibe T, R1…R5 antes do nome do tema (ajuda a distinguir as revisões).")
    item("Grandes áreas", "Edite a lista (até 8). Ela alimenta a lista suspensa e o Painel.")

    h("REMARCAR UMA REVISÃO")
    item("Data manual", "Na aba Temas, digite a nova data por cima da data calculada. Para voltar ao "
                        "automático, copie a mesma coluna de outra linha (Ctrl+C / Ctrl+V).")

    h("LEGENDA DE CORES")
    for nome, code, sigla, c_bar, c_bg, c_ink, c_done in CATS:
        c = put(ws, f"C{r}", f"  {nome}", font=f(9, True, c_ink), fill_=fill(c_bg), align=AL_L)
        c.border = Border(left=Side(style="thick", color=c_bar))
        put(ws, f"D{r}", CAT_DESC[nome], font=f(9, color=INK2))
        r += 1
    c = put(ws, f"C{r}", "  ✓ T · Cardiopediatria", font=f(9, color=CATS[0][6], strike=True),
            fill_=fill(CATS[0][4]))
    c.border = Border(left=Side(style="thick", color=CATS[0][3]))
    put(ws, f"D{r}", "Riscado com ✓ = etapa concluída", font=f(9, color=INK2))
    r += 1
    c = put(ws, f"C{r}", "  Célula amarela", font=f(9, color=INPUT_INK), fill_=fill(INPUT_FILL))
    put(ws, f"D{r}", "Você preenche. Células brancas são calculadas — não digite nelas "
                     "(exceto para remarcar uma data).", font=f(9, color=INK2))
    r += 1

    h("ABAS")
    for nome, desc in (
        ("Calendário", "Visão do mês (como “Mês” no app). Links no topo levam às outras visões."),
        ("Semana", "Visão semanal com até 15 itens por dia."),
        ("Tabela", "Todos os eventos em ordem de data, com filtros (como “Tabela” no app)."),
        ("Painel", "% concluído, atrasadas, o que fazer hoje e a carga dos próximos 14 dias."),
        ("Temas / Avulsos", "Onde você digita."),
        ("Config", "Intervalos, categorias, áreas e preferências."),
        ("Base", "Motor de cálculo. Não edite nem ordene."),
    ):
        item(nome, desc)

    h("LIMITES E CUIDADOS")
    item("Capacidade", f"{MAX_T} temas e {MAX_A} eventos avulsos.")
    item("Estrutura", "Não insira nem exclua colunas. Para tirar um tema, apague o conteúdo da "
                      "linha (Tema, Área, Data e os ✓).")
    item("Ordenar", "Para ver em outra ordem, use os filtros da aba Tabela — não ordene a aba Base.")

    h("DADOS DE EXEMPLO")
    if exemplo:
        item("Este arquivo", "Já vem com as aulas teóricas que aparecem no print (27/09 a 27/10/2026) "
                             "e o simulado UFCSPA 2026, para você ver o calendário funcionando. "
                             "Para começar do zero, use Agendador_de_Revisoes_em_branco.xlsx.")
        item("Itens do print", "Os itens azuis do print que vêm de aulas anteriores a setembro: "
                               "cadastre o tema com a data da aula original e as revisões entram "
                               "sozinhas.")
    else:
        item("Este arquivo", "Versão em branco. Comece pela aba Temas (linha 6). Exemplo de linha: "
                             "Cardiopediatria · Pediatria · 27/09/2026.")


def fit_width(ws, landscape=True):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def build(path, exemplo):
    wb = Workbook()
    ws_como = wb.active
    ws_como.title = "Como Usar"
    ws_temas = wb.create_sheet("Temas")
    ws_av = wb.create_sheet("Avulsos")
    ws_cal = wb.create_sheet("Calendário")
    ws_sem = wb.create_sheet("Semana")
    ws_tab = wb.create_sheet("Tabela")
    ws_pain = wb.create_sheet("Painel")
    ws_cfg = wb.create_sheet("Config")
    ws_base = wb.create_sheet("Base")

    build_config(ws_cfg)
    build_temas(ws_temas, exemplo)
    build_avulsos(ws_av, exemplo)
    build_base(ws_base)
    build_calendario(ws_cal)
    build_semana(ws_sem)
    build_tabela(ws_tab)
    build_painel(ws_pain)
    build_como_usar(ws_como, exemplo)
    fit_width(ws_como, landscape=False)
    fit_width(ws_temas)
    fit_width(ws_av)
    fit_width(ws_tab, landscape=False)

    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = False
    wb.active = wb.sheetnames.index("Calendário")
    ws_cal.sheet_view.tabSelected = True
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)


if __name__ == "__main__":
    out_dir = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parent
    out_dir.mkdir(parents=True, exist_ok=True)
    build(out_dir / "Agendador_de_Revisoes.xlsx", exemplo=True)
    build(out_dir / "Agendador_de_Revisoes_em_branco.xlsx", exemplo=False)
    print("ok:", out_dir)
